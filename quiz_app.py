import streamlit as st
import os
import json
from docx import Document
from io import BytesIO
import pypdf
from google import genai

# =============================================================================
# 1. CẤU HÌNH TRANG WEB & GEMINI CLIENT
# =============================================================================
st.set_page_config(
    page_title="Hệ Thống Tạo Đề Thi Trắc Nghiệm AI",
    page_icon="📝",
    layout="wide"
)

st.title("📝 Hệ Thống Soạn Đề Thi Trắc Nghiệm Tự Động Cho Giáo Viên")
st.write("Tải lên tài liệu bài giảng (PDF/Word/TXT), AI sẽ tự động đọc nội dung và tạo đề thi trắc nghiệm kèm đáp án chi tiết.")

api_key = st.secrets.get("GEMINI_API_KEY") or os.environ.get("GEMINI_API_KEY")

@st.cache_resource
def get_gemini_client(key):
    if not key:
        return None
    return genai.Client(api_key=key)

client = get_gemini_client(api_key)

# =============================================================================
# 2. HÀM ĐỌC FILE VĂN BẢN (PDF, DOCX, TXT)
# =============================================================================
def extract_text_from_file(uploaded_file):
    text = ""
    file_type = uploaded_file.name.split('.')[-1].lower()
    
    try:
        if file_type == 'txt':
            text = uploaded_file.read().decode("utf-8")
        elif file_type == 'pdf':
            pdf_reader = pypdf.PdfReader(uploaded_file)
            for page in pdf_reader.pages:
                text += page.extract_text() + "\n"
        elif file_type in ['docx', 'doc']:
            doc = Document(uploaded_file)
            for para in doc.paragraphs:
                text += para.text + "\n"
    except Exception as e:
        st.error(f"Lỗi khi đọc file: {e}")
    return text

# =============================================================================
# 3. HÀM TẠO ĐỀ THI BẰNG GEMINI AI (ĐỊNH DẠNG JSON)
# =============================================================================
def generate_quiz_json(document_text, num_questions, difficulty, topic):
    if not client:
        return None, "⚠️ Chưa cấu hình GEMINI_API_KEY."

    prompt = f"""
    Bạn là một chuyên gia ra đề thi giáo dục. Hãy đọc tài liệu dưới đây và tạo ra chính xác {num_questions} câu hỏi trắc nghiệm.
    
    Chủ đề/Môn học: {topic}
    Mức độ khó: {difficulty}
    
    Nội dung tài liệu tham khảo:
    \"\"\"
    {document_text[:10000]}
    \"\"\"
    
    YÊU CẦU ĐẦU RA:
    Trả về kết quả dưới dạng một chuỗi JSON thuần túy (JSON Array), KHÔNG chứa ký tự markdown ```json ở đầu hoặc cuối.
    Cấu trúc mỗi phần tử trong mảng JSON như sau:
    [
      {{
        "id": 1,
        "question": "Nội dung câu hỏi?",
        "options": {{
          "A": "Đáp án A",
          "B": "Đáp án B",
          "C": "Đáp án C",
          "D": "Đáp án D"
        }},
        "correct_answer": "A",
        "explanation": "Giải thích chi tiết tại sao A đúng dựa vào tài liệu."
      }}
    ]
    """

    try:
        response = client.models.generate_content(
            model='gemini-3.6-flash',
            contents=prompt,
        )
        clean_json = response.text.strip().replace("```json", "").replace("```", "")
        quiz_data = json.loads(clean_json)
        return quiz_data, None
    except Exception as e:
        return None, f"Lỗi tạo câu hỏi từ Gemini: {e}"

# =============================================================================
# 4. HÀM XUẤT FILE WORD (.DOCX)
# =============================================================================
def create_word_file(quiz_data, title="ĐỀ THI TRẮC NGHIỆM"):
    doc = Document()
    doc.add_heading(title, level=0)
    
    # Phần 1: Đề thi
    doc.add_heading("PHẦN 1: CÂU HỎI TRẮC NGHIỆM", level=1)
    for q in quiz_data:
        p = doc.add_paragraph()
        p.add_run(f"Câu {q['id']}: {q['question']}").bold = True
        for key, value in q['options'].items():
            doc.add_paragraph(f"  {key}. {value}")
        doc.add_paragraph("")
        
    # Phần 2: Đáp án & Giải thích
    doc.add_page_break()
    doc.add_heading("PHẦN 2: ĐÁP ÁN & GIẢI THÍCH CHI TIẾT", level=1)
    for q in quiz_data:
        p = doc.add_paragraph()
        p.add_run(f"Câu {q['id']}: Đáp án đúng {q['correct_answer']}").bold = True
        doc.add_paragraph(f"Giải thích: {q['explanation']}")
        doc.add_paragraph("")
        
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# =============================================================================
# 5. GIAO DIỆN VÀ ĐIỀU KHIỂN
# =============================================================================
st.sidebar.header("⚙️ Cấu Hình Đề Thi")
uploaded_file = st.sidebar.file_uploader("Tải tài liệu bài giảng (PDF, DOCX, TXT):", type=["pdf", "docx", "txt"])

topic_input = st.sidebar.text_input("Tên môn học / Bài học:", value="Tổng hợp")
num_q = st.sidebar.slider("Số lượng câu hỏi cần tạo:", min_value=3, max_value=20, value=5)
difficulty_level = st.sidebar.selectbox("Độ khó:", ["Nhận biết (Dễ)", "Thông hiểu (Trung bình)", "Vận dụng (Khó)"])

if uploaded_file is not None:
    doc_text = extract_text_from_file(uploaded_file)
    
    if len(doc_text.strip()) < 50:
        st.warning("⚠️ Tài liệu quá ngắn hoặc không đọc được nội dung văn bản. Vui lòng kiểm tra lại file.")
    else:
        st.success(f"Trích xuất thành công {len(doc_text)} ký tự từ tài liệu: **{uploaded_file.name}**")
        
        with st.expander("📄 Xem trước nội dung tài liệu trích xuất", expanded=False):
            st.text_area("Nội dung:", doc_text[:1500] + "...", height=150)
            
        if st.button("🚀 Bắt Đầu Tạo Đề Thi Trắc Nghiệm", type="primary"):
            with st.spinner("Gemini AI đang nghiên cứu tài liệu và soạn câu hỏi..."):
                quiz_data, err = generate_quiz_json(doc_text, num_q, difficulty_level, topic_input)
                if err:
                    st.error(err)
                else:
                    st.session_state['quiz_data'] = quiz_data
                    st.success("✨ Đã tạo thành công đề thi!")

if 'quiz_data' in st.session_state:
    quiz_data = st.session_state['quiz_data']
    st.markdown("---")
    st.header("📋 Xem Trước Đề Thi")
    
    for q in quiz_data:
        st.markdown(f"**Câu {q['id']}: {q['question']}**")
        for key, val in q['options'].items():
            st.write(f"- **{key}.** {val}")
        
        with st.expander(f"🔑 Xem đáp án Câu {q['id']}"):
            st.write(f"**Đáp án đúng:** `{q['correct_answer']}`")
            st.write(f"**Giải thích:** {q['explanation']}")
        st.write("")
        
    st.markdown("---")
    col1, col2 = st.columns(2)
    
    with col1:
        docx_file = create_word_file(quiz_data, f"ĐỀ THI: {topic_input.upper()}")
        st.download_button(
            label="📥 Tải Đề Thi File Word (.docx)",
            data=docx_file,
            file_name=f"De_Thi_{topic_input.replace(' ', '_')}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True
        )
        
    with col2:
        json_str = json.dumps(quiz_data, ensure_ascii=False, indent=2)
        st.download_button(
            label="📥 Tải Dữ Liệu Dạng JSON",
            data=json_str,
            file_name="quiz_data.json",
            mime="application/json",
            use_container_width=True
        )